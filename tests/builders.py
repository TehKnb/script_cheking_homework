"""Генератори тестових .docx та зображень воронок."""

from __future__ import annotations

import io
from dataclasses import dataclass, field

import docx
from docx.enum.section import WD_ORIENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image, ImageDraw

from app.visual.layout import Token

RED = RGBColor(0xFF, 0x00, 0x00)
BLUE = RGBColor(0x00, 0x00, 0xFF)
DARK_BLUE = RGBColor(0x15, 0x60, 0x82)  # «Синій, акцент» з теми Word — теж синій
BLACK = RGBColor(0x00, 0x00, 0x00)

SEP = " – "


@dataclass
class Part:
    text: str
    bold: bool = False
    color: RGBColor | None = None
    size: float | None = None
    font: str | None = "Times New Roman"


@dataclass
class DocSpec:
    first_line: list[Part]
    body: list[list[Part]] = field(default_factory=list)
    first_size: float = 14
    body_size: float = 12
    images: list[bytes] = field(default_factory=list)
    hyperlink: str | None = None
    landscape: bool = False


def header_parts(
    name: str = "Іваненко Іван",
    business: str = "Кав'ярня",
    city: str = "Київ",
    package: str = "ВІП67",
    marker: str = "ДЗ №3",
    package_color: RGBColor | None = RED,
) -> list[Part]:
    return [
        Part(name + SEP),
        Part(business, bold=True),
        Part(SEP + city + SEP),
        Part(package, bold=True, color=package_color),
        Part(SEP + marker),
    ]


def plain(*lines: str) -> list[list[Part]]:
    return [[Part(line)] for line in lines]


def build_docx(spec: DocSpec) -> bytes:
    document = docx.Document()
    if spec.landscape:
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    _add_paragraph(document, spec.first_line, spec.first_size)
    for line in spec.body:
        _add_paragraph(document, line, spec.body_size)
    for image in spec.images:
        document.add_picture(io.BytesIO(image))
    if spec.hyperlink:
        _add_hyperlink(document, spec.hyperlink)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_paragraph(document, parts: list[Part], default_size: float) -> None:
    paragraph = document.add_paragraph()
    for part in parts:
        run = paragraph.add_run(part.text)
        run.bold = part.bold
        run.font.size = Pt(part.size or default_size)
        if part.font:
            run.font.name = part.font
        if part.color is not None:
            run.font.color.rgb = part.color


def _add_hyperlink(document, url: str) -> None:
    paragraph = document.add_paragraph()
    r_id = document.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    rpr.extend([fonts, size])
    text = OxmlElement("w:t")
    text.text = url
    run.extend([rpr, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


# ------------------------------------------------------------------ funnels

FUNNEL_W, FUNNEL_H = 1400, 900
# (верх, низ, ліва межа, права межа) кожного блоку
FUNNEL_BLOCKS = [(100, 250, 300, 1100), (250, 400, 400, 1000), (400, 550, 500, 900)]


def funnel_image() -> bytes:
    img = Image.new("RGB", (FUNNEL_W, FUNNEL_H), "white")
    draw = ImageDraw.Draw(img)
    for top, bottom, left, right in FUNNEL_BLOCKS:
        inset = 50
        # Між блоками білий проміжок 4 px — як у піраміді SmartArt
        draw.polygon(
            [(left, top + 2), (right, top + 2), (right - inset, bottom - 2), (left + inset, bottom - 2)],
            fill=(66, 114, 196),
        )
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    return buffer.getvalue()


def blank_image() -> bytes:
    img = Image.new("RGB", (600, 400), (240, 240, 240))
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    return buffer.getvalue()


def tok(text: str, cx: float, cy: float, w: float = 120, h: float = 40, conf: float = 0.95) -> Token:
    return Token(text, cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2, conf)


def funnel_tokens(
    values=("10 000", "2 000", "500"),
    conversions: tuple[str | None, ...] = ("20%", "25%"),
    labels: tuple[str | None, ...] = ("Охоплення", "Ліди", "Продажі"),
) -> list[Token]:
    tokens: list[Token] = []
    for (top, bottom, _, _), value, label in zip(FUNNEL_BLOCKS, values, labels):
        cy = (top + bottom) / 2
        if value:
            tokens.append(tok(value, 700, cy))
        if label:
            tokens.append(tok(label, 1250, cy, w=180))
    boundaries = [FUNNEL_BLOCKS[i][1] for i in range(len(FUNNEL_BLOCKS) - 1)]
    for boundary, conversion in zip(boundaries, conversions):
        if conversion:
            tokens.append(tok(conversion, 150, boundary, w=90))
    return tokens


class FakeOcr:
    name = "fake"

    def __init__(self, tokens: list[Token]) -> None:
        self.tokens = tokens
        self.calls = 0

    def read(self, image) -> list[Token]:
        self.calls += 1
        return list(self.tokens)
