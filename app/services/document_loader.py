"""Завантаження .docx у плоску структуру, зручну для перевірок.

Абзаци розбиваються на рядки по розривах рядка (Shift+Enter): студенти часто пишуть шапку
й точки гіпотези в одному абзаці через розриви, а перевіряти їх треба як окремі рядки.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field, replace
from pathlib import PurePosixPath

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

from app.services.styles import StyleResolver

_RUN_XPATH = (
    "./w:r | ./w:hyperlink/w:r | ./w:ins/w:r | ./w:smartTag/w:r | ./w:customXml/w:r"
    " | ./w:sdt/w:sdtContent/w:r | ./w:fldSimple/w:r | ./w:hyperlink/w:ins/w:r"
)
_HYPERLINK_FIELD_RE = re.compile(r'HYPERLINK\s+"([^"]+)"', re.IGNORECASE)
_A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_EMU_PER_PT = 12700

# Розширення растрових зображень, які можна розпізнати OCR
RASTER_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp"}
# Маркери векторних фігур / SmartArt у document.xml — воронку часто малюють ними
_SHAPE_MARKERS = ("<wps:wsp", "<v:rect", "<v:roundrect", "<v:textbox", "<v:polyline")
_SMARTART_MARKER = "<dgm:relIds"
# Макети SmartArt-пірамід, у яких текст вузлів Word малює всередині блоків (у «Списку-піраміді» — поруч)
_PYRAMID_LAYOUT_RE = re.compile(r"/layout/(?:pyramid1|pyramid3|pyramid4)\b")
_DGM_NS = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
_A_T = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
# Білий текст — службові нотатки, які не видно в документі
_INVISIBLE_COLORS = {"FFFFFF"}
# Кінець шапки: пакет або тип роботи («… – ВІП67 – ДЗ№2»)
_HEADER_TAIL_RE = re.compile(
    r"\bДЗ\s*№?\s*\d|гіпотез|стратегічн|\b(?:віп|vip)\+?\s*\d|\bстандарт\s*\d|\bрічн\w*\s*\d|\bкм\+",
    re.IGNORECASE,
)
# Рядки, які після прийняття дописують співробітники (куратор, гарантійний платіж) — не текст студента
_STAFF_LINE_RE = re.compile(r"^\s*(?:куратор\w*\s*:|гп\b|гарантійн\w*\s+плат)", re.IGNORECASE)


class DocumentReadError(Exception):
    """Файл не є коректним .docx або не читається."""


@dataclass
class RunInfo:
    text: str
    start: int  # зсув у тексті рядка
    end: int
    font_ascii: str
    font_hansi: str
    size_pt: float
    bold: bool
    color: str | None  # HEX або None (auto)
    hyperlink: bool = False


@dataclass
class ParagraphInfo:
    """Один рядок документа (абзац або його частина до розриву рядка)."""

    index: int
    text: str
    runs: list[RunInfo]
    in_table: bool = False

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


@dataclass
class ImageInfo:
    name: str
    data: bytes
    # Розмір, у якому зображення показане в документі (pt); 0 — невідомо
    display_width_pt: float = 0.0
    display_height_pt: float = 0.0

    @property
    def extension(self) -> str:
        return PurePosixPath(self.name).suffix.lower()

    @property
    def is_raster(self) -> bool:
        return self.extension in RASTER_EXTENSIONS


@dataclass
class ParsedDocument:
    paragraphs: list[ParagraphInfo]
    images: list[ImageInfo] = field(default_factory=list)
    shape_count: int = 0
    has_smartart: bool = False
    smartart_without_drawing: bool = False
    hyperlinks: list[str] = field(default_factory=list)
    textbox_text: str = ""
    landscape_sections: int = 0
    # Тексти вузлів SmartArt-пірамід (по одному списку на піраміду): у Word вони стоять усередині блоків
    smartart_pyramids: list[list[str]] = field(default_factory=list)

    @property
    def non_empty_paragraphs(self) -> list[ParagraphInfo]:
        return [p for p in self.paragraphs if not p.is_empty]

    @property
    def first_paragraph(self) -> ParagraphInfo | None:
        return next((p for p in self.paragraphs if not p.is_empty), None)

    @property
    def body_paragraphs(self) -> list[ParagraphInfo]:
        first = self.first_paragraph
        return [
            p
            for p in self.non_empty_paragraphs
            if (first is None or p.index != first.index) and not _STAFF_LINE_RE.match(p.text)
        ]

    @property
    def full_text(self) -> str:
        return "\n".join(p.text for p in self.paragraphs)

    @property
    def body_text(self) -> str:
        return "\n".join(p.text for p in self.body_paragraphs)

    @property
    def has_vector_drawings(self) -> bool:
        """Фігури Word або SmartArt — воронки, оргструктури."""
        return self.has_smartart or self.shape_count > 0

    @property
    def has_vector_funnel_candidate(self) -> bool:
        return self.has_smartart or self.shape_count >= 3


def load_docx(data: bytes, max_unzipped_bytes: int) -> ParsedDocument:
    names = _validate_zip(data, max_unzipped_bytes)
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentReadError(f"python-docx не зміг відкрити файл: {exc}") from exc

    resolver = StyleResolver(document)
    body = document.element.body
    lines: list[ParagraphInfo] = []
    instr_links: list[str] = []
    textbox_parts: list[str] = []

    for p_el in body.iter(qn("w:p")):
        if _inside(p_el, qn("w:txbxContent")):
            # Той самий текстовий блок Word дублює у запасному VML-варіанті (mc:Fallback) — рахуємо один раз
            if _inside(p_el, _MC_FALLBACK):
                continue
            textbox_parts.append("".join(t.text or "" for t in p_el.iter(qn("w:t"))))
            continue
        runs = _read_runs(p_el, resolver)
        for instr in p_el.iter(qn("w:instrText")):
            instr_links.extend(_HYPERLINK_FIELD_RE.findall(instr.text or ""))
        in_table = _inside(p_el, qn("w:tbl"))
        for line_runs in _split_lines(runs):
            lines.append(
                ParagraphInfo(
                    index=len(lines),
                    text="".join(r.text for r in line_runs),
                    runs=line_runs,
                    in_table=in_table,
                )
            )

    xml = document.element.xml
    has_smartart = _SMARTART_MARKER in xml
    return ParsedDocument(
        paragraphs=_merge_header_continuation(lines),
        images=_read_images(data, document),
        shape_count=sum(xml.count(marker) for marker in _SHAPE_MARKERS),
        has_smartart=has_smartart,
        smartart_without_drawing=has_smartart and not any(n.startswith("word/diagrams/drawing") for n in names),
        hyperlinks=_external_hyperlinks(document) + instr_links,
        textbox_text="\n".join(t for t in textbox_parts if t.strip()),
        landscape_sections=sum(
            1 for s in document.sections if s.page_width and s.page_height and s.page_width > s.page_height
        ),
        smartart_pyramids=_read_smartart_pyramids(data, names) if has_smartart else [],
    )


def _merge_header_continuation(lines: list[ParagraphInfo]) -> list[ParagraphInfo]:
    """Довгу шапку Word іноді розбиває на два абзаци: «Іваненко Іван – Юридичний супровід …, -» / «м Вишково,
    Закарпаття – ВІП67 – ДЗ№2». Якщо перший рядок обривається роздільником (або в ньому немає ні пакета,
    ні типу роботи), а наступний короткий рядок того ж розміру шрифту їх містить, — це одна шапка."""
    filled = [line for line in lines if not line.is_empty]
    if len(filled) < 2:
        return lines
    first, nxt = filled[0], filled[1]
    dangling = first.text.rstrip().endswith(("-", "–", "—"))
    if nxt.index - first.index > 2 or len(nxt.text) > 100 or not _HEADER_TAIL_RE.search(nxt.text):
        return lines
    if not dangling and _HEADER_TAIL_RE.search(first.text):
        return lines
    if abs(_main_size(first) - _main_size(nxt)) > 0.5:
        return lines
    shift = len(first.text) + 1
    merged = ParagraphInfo(
        index=first.index,
        text=f"{first.text} {nxt.text}",
        runs=first.runs + [replace(r, start=r.start + shift, end=r.end + shift) for r in nxt.runs],
        in_table=first.in_table,
    )
    return [merged if line is first else line for line in lines if line is not nxt]


def _main_size(line: ParagraphInfo) -> float:
    sizes = sorted(r.size_pt for r in line.runs if r.text.strip())
    return sizes[len(sizes) // 2] if sizes else 0.0


def _read_smartart_pyramids(data: bytes, names: set[str]) -> list[list[str]]:
    """Тексти вузлів SmartArt-пірамід з моделі даних (word/diagrams/dataN.xml). Рендер LibreOffice
    іноді малює піраміду без тексту, а в Word саме ці тексти стоять у блоках."""
    from lxml import etree

    pyramids: list[list[str]] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for name in sorted(n for n in names if re.fullmatch(r"word/diagrams/data\d+\.xml", n)):
            try:
                root = etree.fromstring(zf.read(name))
            except etree.XMLSyntaxError:
                continue
            doc_pt = next((pt for pt in root.iter(f"{_DGM_NS}pt") if pt.get("type") == "doc"), None)
            props = doc_pt.find(f"{_DGM_NS}prSet") if doc_pt is not None else None
            if props is None or not _PYRAMID_LAYOUT_RE.search(props.get("loTypeId") or ""):
                continue
            texts = [
                "".join(t.text or "" for t in pt.iter(_A_T)).strip()
                for pt in root.iter(f"{_DGM_NS}pt")
                if pt.get("type", "node") == "node"
            ]
            pyramids.append([t for t in texts if t])
    return pyramids


# ---------------------------------------------------------------------- helpers


def _validate_zip(data: bytes, max_unzipped_bytes: int) -> set[str]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = set(zf.namelist())
            if "word/document.xml" not in names:
                raise DocumentReadError("У файлі відсутній word/document.xml")
            if sum(info.file_size for info in zf.infolist()) > max_unzipped_bytes:
                raise DocumentReadError("Розпакований розмір документа перевищує допустимий ліміт")
            return names
    except zipfile.BadZipFile as exc:
        raise DocumentReadError("Файл не є коректним ZIP/OOXML-архівом") from exc


def _inside(el, tag: str) -> bool:
    parent = el.getparent()
    while parent is not None:
        if parent.tag == tag:
            return True
        parent = parent.getparent()
    return False


def _run_text(r_el) -> str:
    parts: list[str] = []
    for child in r_el:
        tag = child.tag
        if tag == qn("w:t"):
            parts.append(child.text or "")
        elif tag == qn("w:tab"):
            parts.append("\t")
        elif tag in (qn("w:br"), qn("w:cr")):
            parts.append("\n")
        elif tag == qn("w:noBreakHyphen"):
            parts.append("-")
    return "".join(parts)


def _read_runs(p_el, resolver: StyleResolver) -> list[RunInfo]:
    runs: list[RunInfo] = []
    for r_el in p_el.xpath(_RUN_XPATH):
        text = _run_text(r_el)
        if not text:
            continue
        chain = resolver.rpr_chain(r_el, p_el)
        color = resolver.color(chain)
        if resolver.hidden(chain) or color in _INVISIBLE_COLORS:
            continue
        runs.append(
            RunInfo(
                text=text,
                start=0,
                end=0,
                font_ascii=resolver.font(chain, "ascii"),
                font_hansi=resolver.font(chain, "hAnsi"),
                size_pt=resolver.size_pt(chain),
                bold=resolver.bold(chain),
                color=color,
                hyperlink=r_el.getparent().tag == qn("w:hyperlink")
                or r_el.getparent().getparent().tag == qn("w:hyperlink"),
            )
        )
    return runs


def _split_lines(runs: list[RunInfo]) -> list[list[RunInfo]]:
    """Ділить run-и абзацу на рядки по символу розриву і перераховує зсуви."""
    lines: list[list[RunInfo]] = [[]]
    offset = 0
    for run in runs:
        parts = run.text.split("\n")
        for i, part in enumerate(parts):
            if i > 0:
                lines.append([])
                offset = 0
            if part:
                lines[-1].append(replace(run, text=part, start=offset, end=offset + len(part)))
                offset += len(part)
    return lines


def _read_images(data: bytes, document) -> list[ImageInfo]:
    """Зображення тіла документа через зв'язки (relationships): документи не з Word кладуть їх
    не лише у word/media/."""
    parts: dict[str, object] = {}
    for rel in document.part.rels.values():
        if "image" in rel.reltype and not rel.is_external:
            parts[rel.rId] = rel.target_part

    sizes: dict[str, tuple[float, float]] = {}
    for blip in document.element.body.iter(_A_BLIP):
        if any(a.tag == _MC_FALLBACK for a in blip.iterancestors()):
            continue
        part = parts.get(blip.get(qn("r:embed")))
        holder = next((a for a in blip.iterancestors() if a.tag in (qn("wp:inline"), qn("wp:anchor"))), None)
        extent = holder.find(qn("wp:extent")) if holder is not None else None
        if part is not None and extent is not None:
            sizes[str(part.partname)] = (int(extent.get("cx")) / _EMU_PER_PT, int(extent.get("cy")) / _EMU_PER_PT)

    images: list[ImageInfo] = []
    seen: set[str] = set()
    for part in parts.values():
        name = str(part.partname).lstrip("/")
        if name in seen:
            continue
        seen.add(name)
        w, h = sizes.get("/" + name, (0.0, 0.0))
        images.append(ImageInfo(name=name, data=part.blob, display_width_pt=w, display_height_pt=h))
    return sorted(images, key=lambda i: i.name)


def _external_hyperlinks(document) -> list[str]:
    links: list[str] = []
    for rel in document.part.rels.values():
        if rel.reltype == RT.HYPERLINK and rel.is_external:
            links.append(rel.target_ref)
    return links
